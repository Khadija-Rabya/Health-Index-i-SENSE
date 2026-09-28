"""
Rapport Word : a quel niveau de granularite le Health Index est calcule
(par machine, pas par session).

Reprend l'explication donnee dans la conversation, avec les references de
code exactes (health_index_comparison.py, health_index_pca.py) qui montrent
que l'entrainement PCA/Isolation Forest et les seuils sont calcules une fois
par machine (toutes sessions confondues), et que "session" n'intervient
qu'a posteriori comme test de validation de stabilite.

Sortie : Rapport_Granularite_HealthIndex_iSENSE.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def add_table(doc, headers, rows, caption=None, note=None):
    if caption:
        cap = doc.add_paragraph()
        r = cap.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        cap.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

    if note:
        n = doc.add_paragraph()
        r = n.add_run(note)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        n.paragraph_format.space_before = Pt(4)
        n.paragraph_format.space_after = Pt(12)
    else:
        doc.add_paragraph()
    return table


def bullets(doc, items, style="List Bullet"):
    for it in items:
        doc.add_paragraph(it, style=style)


def numbered(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def paragraph(doc, text, size=11, italic=False, bold=False, color=None, space_after=8):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    r.bold = bold
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def note(doc, text):
    paragraph(doc, text, size=10, italic=True, color=RGBColor(0x40, 0x40, 0x40))


def shade_paragraph(p, fill="F2F2F2"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shd)


def code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        shade_paragraph(p)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Pt(14)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_title_page(doc):
    title = doc.add_heading("Granularité du calcul du Health Index", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title2 = doc.add_heading("Par machine, pas par session", level=1)
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        "À quel niveau le Health Index et ses seuils sont-ils calculés : par session, ou sur "
        "l'ensemble de la machine ? Réponse et références de code exactes."
    )
    run.italic = True
    run.font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        "Projet i-SENSE — Monitoring qualité huile de lubrification\nMotosoufflante A & B — OCP / UM6P"
    ).italic = True
    doc.add_page_break()


def section_reponse(doc):
    doc.add_heading("Réponse directe", level=1)
    p = doc.add_paragraph()
    r = p.add_run(
        "Non — ce n'est pas calculé par session. C'est calculé par machine, sur l'ensemble des "
        "sessions regroupées, et une valeur par ligne/timestamp en sortie."
    )
    r.bold = True
    r.font.size = Pt(12)
    p.paragraph_format.space_after = Pt(10)


def section_ce_qui_est_fait(doc):
    doc.add_heading("Ce qui est réellement fait", level=1)
    paragraph(doc, "Pour chaque machine (A ou B) :")
    numbered(doc, [
        "On prend toutes les lignes de cette machine, toutes sessions confondues (352 sessions "
        "au total sur les 2 machines).",
        "On isole les lignes « saines » parmi ce pool complet (peu importe à quelle session "
        "elles appartiennent).",
        "PCA et Isolation Forest sont entraînés une seule fois par machine sur ce pool de "
        "lignes saines — pas un modèle par session.",
        "Les seuils (percentiles 95/99) sont calculés une seule fois par machine, sur ce même "
        "pool.",
        "Chaque ligne individuelle (peu importe sa session) reçoit ensuite son health_index et "
        "son health_state en comparant ses scores à ces seuils fixes de la machine.",
    ])

    doc.add_heading("Référence de code", level=2)
    paragraph(doc,
        "Dans health_index_comparison.py et health_index_pca.py, la boucle d'entraînement est :"
    )
    code_block(doc, ["for asset, col in ASSET_COLS.items():", "    ..."])
    paragraph(doc,
        "— jamais `for session in sessions`. La colonne session_id n'intervient nulle part dans "
        "l'entraînement ni dans le calcul des seuils."
    )


def section_session(doc):
    doc.add_heading("Où « session » intervient quand même", level=1)
    paragraph(doc,
        "Une seule fois, mais après coup, comme test de validation (Test 3 dans "
        "health_index_comparison.py) : on vérifie que le score ne dérive pas anormalement à "
        "l'intérieur d'une session stable."
    )
    code_block(doc, [
        "ratio = écart-type du score sur une session / écart-type global du score",
        "condition de validation : ratio < 0,7,  calculé sur les 10 sessions les plus",
        "longues de chaque machine",
    ])
    note(doc,
        "C'est un contrôle de qualité, pas une méthode de calcul — la session sert juste à "
        "vérifier a posteriori que le score reste cohérent pendant une période de "
        "fonctionnement continu."
    )


def section_resume(doc):
    doc.add_heading("Résumé", level=1)
    add_table(doc,
        ["Étape", "Granularité"],
        [
            ["Sélection des lignes saines (référence)", "Par machine (toutes sessions confondues)"],
            ["Entraînement PCA / Isolation Forest", "Par machine (1 modèle par machine)"],
            ["Calcul des seuils Surveillance / Alarme", "Par machine (1 jeu de seuils par machine)"],
            ["Valeur finale health_index / health_state", "Par ligne (par timestamp)"],
            ["Rôle de la session", "Validation a posteriori uniquement (stabilité intra-session)"],
        ]
    )
    paragraph(doc,
        "Granularité de calcul = par machine (référence/seuils) → par ligne (valeur finale). "
        "La session est ignorée en entrée, et n'apparaît qu'en sortie comme fenêtre de "
        "contrôle de stabilité."
    )


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_title_page(doc)
    section_reponse(doc)
    section_ce_qui_est_fait(doc)
    section_session(doc)
    section_resume(doc)

    doc.save("Rapport_Granularite_HealthIndex_iSENSE.docx")
    print("Généré : Rapport_Granularite_HealthIndex_iSENSE.docx")


if __name__ == "__main__":
    main()
