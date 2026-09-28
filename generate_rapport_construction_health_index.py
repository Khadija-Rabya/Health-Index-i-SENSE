"""
Rapport Word : comment le Health Index est construit, et comment ses bornes
(Normal / Surveillance / Alarme) sont fixees.

Reprend l'explication donnee au fil de la conversation : principe general
(apprendre le "normal" plutot que la "panne"), construction statistique des
seuils par percentiles sur les lignes saines, logique de combinaison
"le pire gagne" entre PCA et Isolation Forest, et derivation du score continu
health_index a partir des memes seuils.

Sortie : Rapport_Construction_HealthIndex_iSENSE.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def add_table(doc, headers, rows, caption=None, note=None):
    if caption:
        cap = doc.add_paragraph()
        r = cap.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        cap.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

    if note:
        n = doc.add_paragraph()
        r = n.add_run(note)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        n.paragraph_format.space_before = Pt(4)
        n.paragraph_format.space_after = Pt(12)
    else:
        doc.add_paragraph()
    return table


def bullets(doc, items, style="List Bullet"):
    for it in items:
        doc.add_paragraph(it, style=style)


def numbered(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def paragraph(doc, text, size=11, italic=False, bold=False, color=None, space_after=8):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    r.bold = bold
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def note(doc, text):
    paragraph(doc, text, size=10, italic=True, color=RGBColor(0x40, 0x40, 0x40))


def shade_paragraph(p, fill="F2F2F2"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shd)


def formula_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        shade_paragraph(p)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Pt(14)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def flow_step(doc, text):
    p = doc.add_paragraph()
    shade_paragraph(p, fill="E8F0FE")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Pt(10)
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.name = "Consolas"


def add_title_page(doc):
    title = doc.add_heading("Construction du Health Index", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title2 = doc.add_heading("Et de ses bornes Normal / Surveillance / Alarme", level=1)
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        "Comment le score de santé de l'huile est calculé, et comment les seuils qui séparent "
        "un état normal, une zone de surveillance et une alarme sont fixés statistiquement — "
        "à partir des mêmes données, sans historique de pannes."
    )
    run.italic = True
    run.font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        "Projet i-SENSE — Monitoring qualité huile de lubrification\nMotosoufflante A & B — OCP / UM6P"
    ).italic = True
    doc.add_page_break()


def section1(doc):
    doc.add_heading("1. Le principe de base : apprendre ce qu'est « normal », pas ce qu'est « en panne »", level=1)
    paragraph(doc,
        "Il n'existe pas d'historique de pannes réelles dans les données. Impossible donc "
        "d'entraîner un modèle supervisé (« ceci est une panne, ceci n'en est pas une »). La "
        "solution retenue est une approche non-supervisée par référence saine."
    )
    bullets(doc, [
        "Pour chaque machine, on isole un sous-ensemble de lignes « saines » : celles où aucun "
        "flag métier (seuil de la Phase 1) n'est actif.",
        "Deux méthodes indépendantes sont entraînées uniquement sur ces lignes saines, puis "
        "appliquées à toutes les lignes (saines et non saines) : la PCA (Phase 2), qui produit "
        "deux scores par ligne — T² et SPE — et l'Isolation Forest (Phase 3), qui produit un "
        "score d'anomalie par ligne.",
    ])
    paragraph(doc,
        "Chaque méthode apprend ainsi « à quoi ressemble le comportement normal » et mesure "
        "ensuite l'écart de chaque nouvelle ligne à ce comportement."
    )


def section2(doc):
    doc.add_heading("2. Construction des bornes Normal / Surveillance / Alarme", level=1)
    paragraph(doc,
        "C'est ici la logique clé, et elle est statistique, pas arbitraire (convention SPC — "
        "Statistical Process Control)."
    )
    formula_block(doc, [
        "seuil_Surveillance = 95e percentile du score, calculé UNIQUEMENT sur les lignes saines",
        "seuil_Alarme       = 99e percentile du score, calculé UNIQUEMENT sur les lignes saines",
    ])
    paragraph(doc,
        "Concrètement, pour chaque machine et pour chacun des 3 scores (T², SPE, score "
        "Isolation Forest), on prend la distribution de ce score sur les lignes considérées "
        "comme saines, et on regarde où se trouvent ses 95e et 99e centiles."
    )

    doc.add_heading("Pourquoi cette logique ?", level=2)
    paragraph(doc,
        "Même en fonctionnement parfaitement normal, un score de PCA ou d'Isolation Forest "
        "n'est jamais rigoureusement constant — il fluctue. L'idée est : si même les lignes "
        "saines ont un score qui dépasse une valeur dans 5 % des cas, alors dépasser cette "
        "valeur n'est pas forcément anormal — mais dépasser la valeur atteinte seulement 1 % du "
        "temps, ça commence à sortir franchement de la normalité. D'où :"
    )
    add_table(doc,
        ["Zone", "Condition sur le score"],
        [
            ["Normal", "En dessous du 95e percentile des lignes saines"],
            ["Surveillance", "Entre le 95e et le 99e percentile (zone grise, à surveiller)"],
            ["Alarme", "Au-dessus du 99e percentile (rare même pour les lignes saines)"],
        ]
    )
    note(doc,
        "Ces bornes sont recalculées séparément pour chaque machine (A et B ont des seuils "
        "différents) et séparément pour chaque score (T², SPE, Isolation Forest)."
    )


def section3(doc):
    doc.add_heading("3. État par méthode, puis combinaison « le pire gagne »", level=1)
    paragraph(doc,
        "Pour la PCA, l'état d'une ligne est déterminé en comparant ses deux scores bruts "
        "(pas le health_index) à leurs seuils respectifs :"
    )
    formula_block(doc, [
        "état_PCA = Alarme        si T² > T²_alarme  OU  SPE > SPE_alarme",
        "état_PCA = Surveillance  si T² > T²_surv    OU  SPE > SPE_surv   (et pas Alarme)",
        "état_PCA = Normal        sinon",
    ])
    paragraph(doc, "Pour l'Isolation Forest, même logique avec un seul score :")
    formula_block(doc, [
        "état_IsoForest = Alarme / Surveillance / Normal selon score vs iso_surv / iso_alarme",
    ])
    paragraph(doc, "Puis l'état final retenu est le plus sévère des deux méthodes :")
    formula_block(doc, ["état_final = max( état_PCA, état_IsoForest )"])
    paragraph(doc,
        "C'est une logique OR / conservatrice délibérée : si une seule des deux méthodes "
        "détecte une anomalie, l'alerte est levée. Le raisonnement est qu'en système d'alarme "
        "industriel, rater une vraie dégradation coûte plus cher qu'une fausse alerte — donc on "
        "préfère sur-détecter."
    )


def section4(doc):
    doc.add_heading("4. D'où vient le nombre continu health_index (0 à 1)", level=1)
    paragraph(doc,
        "L'état (Normal/Surveillance/Alarme) est une catégorie à 3 niveaux — utile pour un "
        "dashboard, mais grossière pour du machine learning ou du suivi fin. Le health_index "
        "transforme ça en un score continu :"
    )
    formula_block(doc, [
        "sévérité_PCA(x)       = max( T²(x)/seuil_alarme_T² , SPE(x)/seuil_alarme_SPE )",
        "sévérité_IsoForest(x) = score_isolement(x) / seuil_alarme_IsoForest",
        "sévérité_combinée(x)  = max( sévérité_PCA(x), sévérité_IsoForest(x) )",
        "health_index(x)       = 1 / ( 1 + sévérité_combinée(x) )",
    ])
    paragraph(doc,
        "C'est-à-dire : on regarde, pour chaque méthode, à quel pourcentage du seuil d'alarme "
        "se trouve la ligne (ratio, pas juste une comparaison booléenne), on garde le pire des "
        "deux, puis on transforme ce ratio en un nombre borné dans ]0,1] via 1/(1+x)."
    )

    doc.add_heading("Point important à ne pas confondre", level=2)
    paragraph(doc,
        "health_state (la catégorie) et health_index (le nombre) sont calculés par deux "
        "chemins parallèles à partir des mêmes seuils, pas l'un à partir de l'autre :"
    )
    bullets(doc, [
        "health_state = comparaison discrète des scores bruts aux seuils Surveillance/Alarme",
        "health_index = transformation continue du ratio « score / seuil d'Alarme »",
    ])
    paragraph(doc,
        "Ils sont cohérents entre eux à un point précis : quand la ligne est exactement au "
        "seuil d'Alarme sur le pire indicateur, sévérité_combinée = 1, donc health_index = 0,5. "
        "Mais il n'y a pas de valeur unique de health_index qui corresponde exactement à la "
        "frontière Surveillance, car cette frontière dépend du ratio "
        "seuil_surveillance/seuil_alarme, qui varie selon la machine et le score — c'est pour "
        "ça qu'on garde les deux colonnes (health_index continu + health_state catégoriel) "
        "plutôt que de redériver l'état en découpant health_index à des bornes fixes."
    )


def section5(doc):
    doc.add_heading("5. Résumé du flux complet", level=1)
    flow_step(doc, "Lignes saines (Phase 1)")
    flow_step(doc, "  │")
    flow_step(doc, "  ├──> PCA entraînée dessus ──> T²(x), SPE(x) pour toutes les lignes")
    flow_step(doc, "  │         │")
    flow_step(doc, "  │         └──> seuils = percentiles 95/99 de T², SPE sur les lignes saines")
    flow_step(doc, "  │")
    flow_step(doc, "  └──> Isolation Forest entraînée dessus ──> score(x) pour toutes les lignes")
    flow_step(doc, "            │")
    flow_step(doc, "            └──> seuils = percentiles 95/99 du score sur les lignes saines")
    doc.add_paragraph()
    flow_step(doc, "état_PCA (Normal/Surv/Alarme)  +  état_IsoForest (Normal/Surv/Alarme)")
    flow_step(doc, "                    │")
    flow_step(doc, "              max (le pire gagne) ──> health_state final")
    flow_step(doc, "                    │")
    flow_step(doc, "sévérité_PCA, sévérité_IsoForest ──> max ──> health_index = 1/(1+sévérité)")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_title_page(doc)
    section1(doc)
    section2(doc)
    section3(doc)
    section4(doc)
    section5(doc)

    doc.save("Rapport_Construction_HealthIndex_iSENSE.docx")
    print("Généré : Rapport_Construction_HealthIndex_iSENSE.docx")


if __name__ == "__main__":
    main()
