"""
Rapport Word : formules utilisees a chaque phase de construction du Health
Index (explication + resultats), et tableau complet des metriques de
prediction (tous les modeles, pas seulement le meilleur).

Reprend le contenu de explication_formules_health_index.md, mis en forme en
Word avec les memes conventions que les autres rapports du projet.

Sortie : Rapport_Formules_HealthIndex_iSENSE.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def add_table(doc, headers, rows, caption=None, note=None):
    if caption:
        cap = doc.add_paragraph()
        r = cap.add_run(caption)
        r.italic = True
        r.font.size = Pt(9)
        cap.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)

    if note:
        n = doc.add_paragraph()
        r = n.add_run(note)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        n.paragraph_format.space_before = Pt(4)
        n.paragraph_format.space_after = Pt(12)
    else:
        doc.add_paragraph()
    return table


def bullets(doc, items, style="List Bullet"):
    for it in items:
        doc.add_paragraph(it, style=style)


def paragraph(doc, text, size=11, italic=False, bold=False, color=None, space_after=8):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    r.bold = bold
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def note(doc, text):
    paragraph(doc, text, size=10, italic=True, color=RGBColor(0x40, 0x40, 0x40))


def shade_paragraph(p, fill="F2F2F2"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p._p.get_or_add_pPr().append(shd)


def formula_block(doc, lines):
    """Bloc de formule : police monospace, fond gris clair, une ligne par element de `lines`."""
    for line in lines:
        p = doc.add_paragraph()
        shade_paragraph(p)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Pt(14)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_title_page(doc):
    title = doc.add_heading("Formules du Health Index", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title2 = doc.add_heading("Détail par phase, variables, et résultats", level=1)
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        "Pour chaque phase : les variables utilisées, la formule mathématique exacte du code, "
        "une explication de ce qui se passe, et les résultats chiffrés obtenus. Complété par le "
        "tableau intégral des métriques de tous les modèles de prédiction."
    )
    run.italic = True
    run.font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        "Projet i-SENSE — Monitoring qualité huile de lubrification\nMotosoufflante A & B — OCP / UM6P"
    ).italic = True
    doc.add_page_break()


# ===========================================================================
# PHASE 1
# ===========================================================================

def phase1(doc):
    doc.add_heading("Phase 1 — Règles métier à seuils officiels", level=1)
    paragraph(doc,
        "12 variables santé sont chacune converties en une distance de dégradation entre 0 (état "
        "sain) et 1 (seuil « Surveillance » officiel OCP atteint), puis combinées en un score "
        "unique. Il y a 3 formes de distance selon la nature de la variable."
    )

    doc.add_heading("Variables utilisées (12) et pourquoi", level=2)
    add_table(doc,
        ["Variable", "Colonne réelle", "Formule utilisée"],
        [
            ["Oil Temperature", "Oil Temperature_filled", "Unidirectionnelle"],
            ["Oil H2O Saturation", "Oil H2O Saturation_filled", "Unidirectionnelle"],
            ["DC", "DC_filled", "Unidirectionnelle (élevée = critique)"],
            ["Oil Conductivity", "Oil Conductivity_nSm_filled", "Bidirectionnelle"],
            ["ISO 4", "ISO 4_filled", "Unidirectionnelle"],
            ["ISO 6", "ISO 6", "Unidirectionnelle"],
            ["ISO 14", "ISO 14", "Unidirectionnelle"],
            ["Oil H2O ppm", "Oil H2O ppm_filled", "Unidirectionnelle"],
            ["Density", "Density_filled", "Écart % (référence 858 kg/m³)"],
            ["Kinematic Viscosity", "Kinematic Viscosity_filled", "Écart % (référence = médiane machine)"],
            ["Dynamic Viscosity", "Dynamic Viscosity_filled", "Écart % (référence = médiane machine)"],
            ["Viscosity at 40°C", "Viscosity at 40°C_filled", "Écart % (référence 46,0 cSt)"],
        ],
        note="Ce sont toutes les variables pour lesquelles le document officiel fournit des "
             "seuils — aucune n'est exclue ici. Oil System Vibration n'a pas de seuil officiel "
             "et n'est donc pas incluse (elle l'est en Phase 2)."
    )

    doc.add_heading("Formule 1 — Variables unidirectionnelles (« plus haut = pire »)", level=2)
    formula_block(doc, ["distance(x) = clip( (x - seuil_normal) / (seuil_surveillance - seuil_normal),  0, 1 )"])
    paragraph(doc,
        "x = valeur mesurée ; seuil_normal / seuil_surveillance = bornes officielles OCP ; "
        "clip(v, 0, 1) ramène v dans [0, 1]. Si x est sous le seuil normal, distance = 0 ; "
        "au-delà du seuil de surveillance, distance reste plafonnée à 1."
    )

    doc.add_heading("Formule 2 — Variable bidirectionnelle (Oil Conductivity_nSm)", level=2)
    formula_block(doc, [
        "distance_bas(x)  = (seuil_normal_bas - x) / (seuil_normal_bas - seuil_surveillance_bas)",
        "distance_haut(x) = (x - seuil_normal_haut) / (seuil_surveillance_haut - seuil_normal_haut)",
        "distance(x) = clip( max(distance_bas(x), distance_haut(x), 0),  0, 1 )",
    ])
    paragraph(doc,
        "La conductivité a une plage « Normal » à deux bornes (0,4-2,0 nS/m) : une valeur trop "
        "basse ou trop haute est anormale. Dans la plage normale, les deux termes sont négatifs "
        "et le max(..., 0) ramène la distance à 0."
    )

    doc.add_heading("Formule 3 — Écart relatif à une référence (densité, viscosités)", level=2)
    formula_block(doc, [
        "écart_% = | x - référence | / référence × 100",
        "distance(x) = clip( (écart_% - normal_%) / (surveillance_% - normal_%),  0, 1 )",
    ])
    paragraph(doc,
        "référence = 858 kg/m³ (Density) ou 46,0 cSt (Viscosity at 40°C), valeurs TD46/Shell "
        "Turbo T46 ; médiane de la machine pour Kinematic/Dynamic Viscosity, faute de référence "
        "absolue officielle."
    )

    doc.add_heading("Formule de combinaison finale", level=2)
    formula_block(doc, ["HI_regles = clip( 1 - (1/12) × Σ(distance_i pour i=1..12),  0, 1 )"])
    paragraph(doc, "Moyenne simple des 12 distances (poids égal, faute de pondération officielle connue).")

    doc.add_heading("Résultats obtenus", level=2)
    add_table(doc,
        ["Machine", "HI moyen", "HI min", "% Surveillance", "% Alarme"],
        [
            ["Motosoufflante A", "0,814", "0,482", "3,9 %", "1,5 %"],
            ["Motosoufflante B", "0,606", "0,450", "6,0 %", "1,3 %"],
        ]
    )
    add_table(doc,
        ["État", "Lignes", "%"],
        [
            ["Normal", "39 421", "93,6 %"],
            ["Surveillance", "2 123", "5,0 %"],
            ["Alarme", "597", "1,4 %"],
        ]
    )


# ===========================================================================
# PHASE 2
# ===========================================================================

def phase2(doc):
    doc.add_page_break()
    doc.add_heading("Phase 2 — PCA par machine, T² et SPE", level=1)
    paragraph(doc,
        "13 variables santé (DC exclue) sont standardisées puis projetées sur un sous-espace de "
        "5 composantes principales, appris uniquement sur les lignes saines (sans flag actif). "
        "Deux statistiques mesurent ensuite l'écart de chaque nouvelle ligne à ce sous-espace."
    )

    doc.add_heading("Variables utilisées (13) — et ce que sont vraiment les « 5 composantes »", level=2)
    note(doc,
        "Point important : les « 5 composantes retenues » ne sont PAS 5 des 13 variables "
        "sélectionnées parmi les autres — la PCA garde les 13 variables, mais les recombine "
        "mathématiquement en 5 nouvelles variables synthétiques (combinaisons linéaires "
        "pondérées des 13 originales), qui suffisent à elles seules à reconstituer 94 % de "
        "l'information. Aucune des 13 variables ci-dessous n'est individuellement supprimée par "
        "la PCA — c'est le choix de la liste de 13 (avant la PCA) qui exclut certaines "
        "variables du dataset complet."
    )
    add_table(doc,
        ["#", "Variable", "Colonne réelle"],
        [
            ["1", "ISO 4", "ISO 4_filled"],
            ["2", "ISO 6", "ISO 6"],
            ["3", "ISO 14", "ISO 14"],
            ["4", "Oil H2O ppm", "Oil H2O ppm_filled"],
            ["5", "Oil H2O Saturation", "Oil H2O Saturation_filled"],
            ["6", "Viscosity at 40°C", "Viscosity at 40°C_filled"],
            ["7", "Kinematic Viscosity", "Kinematic Viscosity_filled"],
            ["8", "Écart au grade de viscosité", "viscosity_grade_gap"],
            ["9", "Proxy Indice de Viscosité", "vi_proxy"],
            ["10", "Densité normalisée 15°C", "density_15C"],
            ["11", "Écart de densité", "density_deviation_pct"],
            ["12", "Oil System Vibration", "Oil System Vibration_filled"],
            ["13", "Oil Temperature", "Oil Temperature_filled"],
        ]
    )

    doc.add_heading("Pourquoi seulement ces 13, sur les ~150 colonnes du dataset complet", level=2)
    add_table(doc,
        ["Variable(s) écartée(s)", "Raison"],
        [
            ["DC / DC_filled",
             "Exclue explicitement : 88,4 % des mesures de Motosoufflante B sont physiquement "
             "impossibles (anomalie de calibration, cf. Phase 0). L'inclure injecterait un "
             "artefact de capteur plutôt qu'un vrai signal."],
            ["Oil Conductivity_nSm_filled",
             "Non incluse ici (elle l'est en Phase 1) — liste PCA non retouchée lors de l'ajout "
             "de la conductivité au pipeline."],
            ["Lags, EWMA, tendance, z-score (~100 colonnes)",
             "Transformations temporelles des mêmes variables déjà présentes sous forme "
             "instantanée — redondance massive (quasi colinéarité) sans information nouvelle."],
            ["Calendaires, session",
             "Décrivent le contexte temporel, pas l'état physico-chimique de l'huile — hors "
             "périmètre d'un indice de santé."],
            ["Flags de seuils, confiance vibration/capteur",
             "Métadonnées dérivées des variables continues déjà incluses — dupliqueraient leur "
             "information sous une autre forme."],
            ["Density_filled, Dynamic Viscosity_filled (brutes)",
             "Remplacées par leurs versions transformées plus informatives (density_15C/"
             "density_deviation_pct, Kinematic Viscosity déjà présente)."],
        ]
    )

    doc.add_heading("Formule 0 — Standardisation", level=2)
    formula_block(doc, ["x_standardisé = (x - moyenne_entraînement) / écart_type_entraînement"])
    paragraph(doc, "Centrage-réduction avant la PCA, pour que les variables à grande échelle "
                   "(ex. ISO 4406) ne dominent pas artificiellement celles à petite échelle "
                   "(ex. densité, quelques kg/m³ de variation).")

    doc.add_heading("Formule 1 — T² de Hotelling", level=2)
    formula_block(doc, ["T²(x) = Σ ( score_i(x)² / valeur_propre_i )   pour i = 1..5"])
    paragraph(doc,
        "score_i(x) = coordonnée de x sur la i-ème composante principale ; valeur_propre_i = "
        "variance expliquée par cette composante. Distance de Mahalanobis dans l'espace réduit "
        "des 5 dimensions de variation normale."
    )

    doc.add_heading("Formule 2 — SPE (erreur de reconstruction)", level=2)
    formula_block(doc, [
        "x̂ = reconstruction de x_standardisé à partir des 5 scores (projection inverse)",
        "SPE(x) = Σ ( x_standardisé - x̂ )²",
    ])
    paragraph(doc,
        "Mesure ce que les 5 composantes ne parviennent PAS à expliquer de x — signal différent "
        "de T² (qui reste dans la structure apprise)."
    )

    doc.add_heading("Seuils", level=2)
    formula_block(doc, [
        "seuil_surveillance = 95e percentile de T² (ou SPE), calculé sur les lignes SAINES uniquement",
        "seuil_alarme        = 99e percentile de T² (ou SPE), calculé sur les lignes SAINES uniquement",
    ])

    doc.add_heading("Résultats obtenus", level=2)
    add_table(doc,
        ["Machine", "Composantes (k)", "Variance expliquée", "Seuil Surv. T²", "Seuil Alarme T²", "Seuil Surv. SPE", "Seuil Alarme SPE"],
        [
            ["A", "5", "94,3 %", "15,10", "29,72", "4,10", "8,46"],
            ["B", "5", "94,4 %", "12,87", "23,57", "2,47", "5,13"],
        ]
    )
    add_table(doc,
        ["État (PCA)", "Lignes"],
        [
            ["Normal", "37 888"],
            ["Surveillance", "3 085"],
            ["Alarme", "1 168"],
        ]
    )


# ===========================================================================
# PHASE 3
# ===========================================================================

def phase3(doc):
    doc.add_page_break()
    doc.add_heading("Phase 3 — Isolation Forest", level=1)
    paragraph(doc,
        "Exactement les mêmes 13 variables que la Phase 2 (DC toujours exclue pour la même "
        "raison), mêmes lignes saines d'entraînement, mais avec une méthode qui ne suppose ni "
        "linéarité ni distribution gaussienne — vérification croisée de la PCA. Pas de réduction "
        "à un nombre restreint de composantes : travaille directement sur les 13 dimensions."
    )

    doc.add_heading("Formule — score d'isolement", level=2)
    formula_block(doc, [
        "E[h(x)] = profondeur moyenne d'isolement de x, sur tous les arbres de la forêt",
        "c(n)     = profondeur moyenne attendue pour un point isolé par hasard (taille n)",
        "score_isolement(x) = 2^( - E[h(x)] / c(n) )",
    ])
    paragraph(doc,
        "Un point normal (cœur de la masse des données) demande beaucoup de coupures pour être "
        "isolé (E[h(x)] grand → score proche de 0) ; un point anormal (périphérie) est isolé en "
        "très peu de coupures (score proche de 1). Le code utilise `-score_samples(x)` "
        "(scikit-learn), même convention que T²/SPE : plus haut = plus anormal."
    )

    doc.add_heading("Résultats obtenus", level=2)
    add_table(doc,
        ["Machine", "Seuil Surveillance", "Seuil Alarme"],
        [
            ["A", "0,606", "0,632"],
            ["B", "0,577", "0,643"],
        ]
    )
    add_table(doc,
        ["État (Isolation Forest)", "Lignes"],
        [
            ["Normal", "39 659"],
            ["Surveillance", "2 001"],
            ["Alarme", "481"],
        ],
        note="Cohérence avec la PCA : accord parfait sur la zone Normal (0 ligne classée Normal "
             "par la PCA et Alarme par l'Isolation Forest)."
    )


# ===========================================================================
# PHASE 4 & 5
# ===========================================================================

def phase4_5(doc):
    doc.add_page_break()
    doc.add_heading("Phase 4 — Décision (pas de formule)", level=1)
    paragraph(doc,
        "3 tests statistiques (cohérence avec les flags métier, injection de défauts "
        "synthétiques, stabilité en session stable) ont validé la PCA/Isolation Forest sans "
        "réserve — l'autoencodeur n'a donc pas été construit."
    )

    doc.add_heading("Phase 5 — Health Index final", level=1)
    paragraph(doc,
        "Les scores de la PCA (T², SPE) et de l'Isolation Forest sont combinés en un seul score "
        "continu, selon une logique conservatrice : le pire des deux modèles l'emporte."
    )
    note(doc,
        "Variables utilisées : aucune nouvelle variable brute — cette phase combine uniquement "
        "les scores déjà calculés en Phases 2 et 3 (T², SPE, score Isolation Forest)."
    )

    doc.add_heading("Formule — sévérité et score final", level=2)
    formula_block(doc, [
        "sévérité_PCA(x)       = max( T²(x)/seuil_alarme_T² , SPE(x)/seuil_alarme_SPE )",
        "sévérité_IsoForest(x) = score_isolement(x) / seuil_alarme_IsoForest",
        "sévérité_combinée(x)  = max( sévérité_PCA(x) , sévérité_IsoForest(x) )",
        "health_index(x)       = 1 / ( 1 + sévérité_combinée(x) )",
    ])
    bullets(doc, [
        "Bornée entre 0 et 1 : sévérité_combinée ≥ 0 donc health_index ∈ ]0, 1].",
        "Vaut exactement 0,5 au seuil d'alarme : quand sévérité_combinée = 1, health_index = 1/(1+1) = 0,5.",
        "Décroissance douce, jamais de chute brutale à 0 : même sévérité=10 donne encore health_index ≈ 0,09.",
    ])

    doc.add_heading("Résultats obtenus", level=2)
    add_table(doc,
        ["Machine", "HI moyen", "HI min", "Normal", "Surveillance", "Alarme"],
        [
            ["A", "0,575", "0,188", "17 658", "1 421", "817"],
            ["B", "0,583", "0,043", "19 928", "1 698", "619"],
            ["Total", "—", "—", "37 586 (89,2 %)", "3 119 (7,4 %)", "1 436 (3,4 %)"],
        ]
    )


# ===========================================================================
# PREDICTION — TOUTES LES METRIQUES
# ===========================================================================

def prediction_metrics(doc):
    doc.add_page_break()
    doc.add_heading("Prédiction du Health Index à t+n — toutes les métriques, tous les modèles", level=1)
    paragraph(doc,
        "Contrairement aux tableaux « meilleur modèle par cas », voici l'intégralité des "
        "résultats (MAE, RMSE, R²) pour les 5 modèles entraînés, sur les 3 horizons communs. "
        "Formulation en delta (health_index[t+n] - health_index[t]), découpage temporel par "
        "machine."
    )

    def horizon_table(title, rows):
        doc.add_heading(title, level=2)
        add_table(doc, ["Machine", "Modèle", "MAE", "RMSE", "R²"], rows)

    horizon_table("10 minutes", [
        ["A", "Random Forest", "0,0037", "0,0071", "0,9453"],
        ["A", "XGBoost", "0,0038", "0,0072", "0,9445"],
        ["A", "A. Autoencodeur+RF (latent)", "0,0038", "0,0073", "0,9422"],
        ["A", "B. LSTM Encoder-Decoder", "0,0081", "0,0151", "0,7557"],
        ["A", "C. Débruiteur + fine-tuning", "0,0063", "0,0111", "0,8683"],
        ["B", "Random Forest", "0,0120", "0,0204", "0,7053"],
        ["B", "XGBoost", "0,0438", "0,0657", "−2,0505"],
        ["B", "A. Autoencodeur+RF (latent)", "0,0129", "0,0226", "0,6395"],
        ["B", "B. LSTM Encoder-Decoder", "0,0211", "0,0321", "0,2696"],
        ["B", "C. Débruiteur + fine-tuning", "0,0135", "0,0244", "0,5776"],
    ])

    horizon_table("3 heures", [
        ["A", "Random Forest", "0,0092", "0,0174", "0,4994"],
        ["A", "XGBoost", "0,0204", "0,0367", "−1,2293"],
        ["A", "A. Autoencodeur+RF (latent)", "0,0074", "0,0152", "0,6172"],
        ["A", "B. LSTM Encoder-Decoder", "0,0595", "0,0773", "−8,8734"],
        ["A", "C. Débruiteur + fine-tuning", "0,0137", "0,0211", "0,2647"],
        ["B", "Random Forest", "0,0150", "0,0248", "0,5572"],
        ["B", "XGBoost", "0,0179", "0,0263", "0,5029"],
        ["B", "A. Autoencodeur+RF (latent)", "0,0160", "0,0277", "0,4482"],
        ["B", "B. LSTM Encoder-Decoder", "0,0196", "0,0285", "0,4126"],
        ["B", "C. Débruiteur + fine-tuning", "0,0239", "0,0522", "−0,9680"],
    ])

    horizon_table("24 heures", [
        ["A", "Random Forest", "0,0278", "0,0342", "−1,3344"],
        ["A", "XGBoost", "0,0189", "0,0263", "−0,3761"],
        ["A", "A. Autoencodeur+RF (latent)", "0,0139", "0,0239", "−0,1406"],
        ["A", "B. LSTM Encoder-Decoder", "0,0152", "0,0235", "−0,1027"],
        ["A", "C. Débruiteur + fine-tuning", "0,0126", "0,0204", "0,1720"],
        ["B", "Random Forest", "0,0371", "0,0422", "−0,1194"],
        ["B", "XGBoost", "0,0448", "0,0528", "−0,7591"],
        ["B", "A. Autoencodeur+RF (latent)", "0,0278", "0,0415", "−0,0826"],
        ["B", "B. LSTM Encoder-Decoder", "0,0634", "0,0734", "−2,3916"],
        ["B", "C. Débruiteur + fine-tuning", "0,0318", "0,0479", "−0,4423"],
    ])

    doc.add_heading("Horizons supplémentaires (Random Forest / XGBoost uniquement)", level=2)
    add_table(doc,
        ["Horizon", "Machine", "Modèle", "MAE", "RMSE", "R²"],
        [
            ["1h", "A", "Random Forest", "0,0059", "0,0127", "0,7847"],
            ["1h", "A", "XGBoost", "0,0062", "0,0117", "0,8194"],
            ["1h", "B", "Random Forest", "0,0221", "0,0281", "0,4186"],
            ["1h", "B", "XGBoost", "0,0150", "0,0235", "0,5917"],
            ["3j", "A", "Random Forest", "0,0211", "0,0339", "−1,4056"],
            ["3j", "A", "XGBoost", "0,0405", "0,0512", "−4,4663"],
            ["3j", "B", "Random Forest", "0,0386", "0,0478", "−0,1864"],
            ["3j", "B", "XGBoost", "0,1177", "0,1389", "−9,0217"],
            ["1 semaine", "A", "Random Forest", "0,0329", "0,0394", "−2,4695"],
            ["1 semaine", "A", "XGBoost", "0,0321", "0,0419", "−2,9268"],
            ["1 semaine", "B", "Random Forest", "0,0199", "0,0280", "−1,1111"],
            ["1 semaine", "B", "XGBoost", "0,0310", "0,0489", "−5,4415"],
        ],
        note="Non testés sur les 3 modèles autoencodeur (limitation de temps de calcul, cf. "
             "plan_prediction_health_index.md)."
    )

    doc.add_heading("Classement global (rang moyen de R², 3 horizons communs)", level=2)
    add_table(doc,
        ["Rang", "Modèle", "Rang moyen", "Meilleur (/6)"],
        [
            ["1", "Random Forest", "2,00", "3"],
            ["2", "A. Autoencodeur+RF (latent)", "2,17", "2"],
            ["3", "C. Débruiteur + fine-tuning", "3,17", "1"],
            ["4", "XGBoost", "3,50", "0"],
            ["5", "B. LSTM Encoder-Decoder", "4,17", "0"],
        ]
    )


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_title_page(doc)
    phase1(doc)
    phase2(doc)
    phase3(doc)
    phase4_5(doc)
    prediction_metrics(doc)

    doc.save("Rapport_Formules_HealthIndex_iSENSE.docx")
    print("Généré : Rapport_Formules_HealthIndex_iSENSE.docx")


if __name__ == "__main__":
    main()
